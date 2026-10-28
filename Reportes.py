import streamlit as st
import pandas as pd
import numpy as np

from io import BytesIO
import openpyxl
import base64

from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# ==========================================
# CONFIGURACIÓN DE LA PÁGINA
# ==========================================
st.set_page_config(page_title="Gestión de Rotables", layout="wide")

@st.cache_data
def cargar_excel(file,tipo,nombre_hoja):
    # 1. Leer el lienzo de la hoja específica
        file.seek(0)
        df_raw = pd.read_excel(file, sheet_name=nombre_hoja, header=None, dtype=str)
        
        # 2. Escanear fila por fila buscando los encabezados reales
        fila_header = 0
        for idx, row in df_raw.iterrows():
            fila_texto = str(row.values).upper()
            if tipo == "plantilla" and ("NUMERO DE SERIE" in fila_texto or "DESCRIPCIÓN" in fila_texto):
                fila_header = idx
                break
            elif tipo == "sistema" and ("SN" in fila_texto or "PN" in fila_texto):
                fila_header = idx
                break
                
        # 3. Reiniciar el archivo y leerlo oficialmente desde la fila encontrada
        file.seek(0)
        df = pd.read_excel(file, sheet_name=nombre_hoja, header=fila_header, dtype=str)
        
        # 4. Limpieza profunda: Quitar columnas y filas fantasma
        df = df.loc[:, ~df.columns.str.contains('^Unnamed')]
        df.columns = df.columns.str.strip().str.upper()
        df = df.dropna(how='all')

        df = df.replace('_x0000_', '', regex=True)
        df = df.replace('\n','',regex=True)
        
        return df

def load_icon(ruta_imagen):
    with open(ruta_imagen,"rb") as image_file:
        encoded_string = base64.b64encode(image_file.read()).decode()
    return f"data:image/png;base64,{encoded_string}"

icono_reporte = load_icon("icons/reporte-de-negocios.png")
icono_alerta = load_icon("icons/peligro.png")
icono_cruce = load_icon("icons/cruce-de-caminos.png")
# ==========================================
# SISTEMA DE AUTENTICACIÓN GLOBAL
# ==========================================
def check_password():
    if "password_correct" not in st.session_state:
        st.session_state["password_correct"] = False

    if not st.session_state["password_correct"]:
        st.title("Acceso al Sistema")
        st.write("Por favor, ingresa la contraseña para continuar.")
        pwd = st.text_input("Contraseña:", type="password")
        
        if pwd == st.secrets.get("admin_password", ""):
            st.session_state["password_correct"] = True
            st.rerun()
        elif pwd != "":
            st.error("Contraseña incorrecta.")
        return False
    return True

# ==========================================
# APLICACIÓN PRINCIPAL
# ==========================================
if check_password():
    banner_html = """
    <style>
        .hero-banner {
            background: linear-gradient(135deg, #490359 0%, #E607E9 100%);
            padding: 2.5rem 2rem;
            border-radius: 12px;
            margin-bottom: 2rem;
            box-shadow: 0 6px 10px rgba(0, 0, 0, 0.15);
            color: white;
            text-align: left;
            display: flex;
            flex-direction: column;
            justify-content: center;
        }
        .hero-title {
            font-size: 2.4rem;
            font-weight: 800;
            margin: 0;
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            letter-spacing: 0.5px;
        }
        .hero-subtitle {
            font-size: 1.1rem;
            font-weight: 400;
            margin-top: 8px;
            opacity: 0.9;
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
        }
    </style>

    <div class="hero-banner">
        <div class="hero-title">SISTEMA DE REPORTES DE ROTABLES</div>
    </div>
    """

    st.markdown(banner_html, unsafe_allow_html=True)

    st.markdown("""
    <style>
        button[data-baseweb="tab"] p, 
        button[data-baseweb="tab"] span, 
        div[data-testid="stTabs"] button p {
            font-size: 22px !important; 
            font-weight: 800 !important;
            text-transform: uppercase !important;
        }

        /* Cambiar la línea roja por azul oscuro */
        button[data-baseweb="tab"][aria-selected="true"],
        div[data-testid="stTabs"] button[aria-selected="true"] {
            border-bottom: 4px solid #002060 !important;
        }

        /* Cambiar el color de la letra activa a azul */
        button[data-baseweb="tab"][aria-selected="true"] p,
        button[data-baseweb="tab"][aria-selected="true"] span {
            color: #002060 !important;
        }

        /* Efecto Hover */
        button[data-baseweb="tab"]:hover p {
            color: #17a2b8 !important;
        }

        /* ----------------------------------------------------
           2. ESTILOS PARA LAS CAJAS DE SUBIDA DE ARCHIVOS
           ---------------------------------------------------- */
        
        /* Caja 1 (Azul) - Apuntando directamente a la etiqueta 'section' */
        [data-testid="stColumn"]:nth-of-type(1) section[data-testid="stFileUploadDropzone"] {
            background-color: #072a40 !important;
            border: 2px dashed #17a2b8 !important;
            border-radius: 10px !important;
        }
        
        /* Caja 2 (Verde) */
        [data-testid="stColumn"]:nth-of-type(2) section[data-testid="stFileUploadDropzone"] {
            background-color: #053015 !important;
            border: 2px dashed #28a745 !important;
            border-radius: 10px !important;
        }

        /* Asegurar que el texto dentro de las cajas oscuras sea blanco */
        section[data-testid="stFileUploadDropzone"] * {
            color: white !important;
        }
        
        /* Darle estilo al botoncito de "Browse files" para que no se pierda */
        section[data-testid="stFileUploadDropzone"] button {
            background-color: white !important;
            color: #002060 !important;
            font-weight: bold !important;
        }
        div[data-testid="stVerticalBlock"]:has(> div.element-container .marcador-alerta) {
            background-color: #fdf3f4 !important; /* Fondo rojito/rosado muy suave */
            padding: 25px !important;
            border-radius: 12px !important;
            border: 2px solid #f8d7da !important; /* Borde rojo pastel */
            margin-top: 20px !important;
            margin-bottom: 20px !important;
        }
        div[data-testid="stVerticalBlock"]:has(> div.element-container .marcador-cruce) {
            background-color: #FCDEFF !important; 
            border: 2px solid #F586FF !important;
            padding: 25px !important;
            border-radius: 12px !important;
            margin-top: 20px !important;
            margin-bottom: 20px !important;
        }
        div[data-testid="stVerticalBlock"]:has(> div.element-container .marcador-tabla) {
                    background-color: #EDEDED!important; 
                    border: 2px solid #C0C0C0!important;
                    padding: 10px !important;
                    border-radius: 12px !important;
                    margin-top: 20px !important;
                    margin-bottom: 20px !important;
                } 
    </style>
    """, unsafe_allow_html=True)
    tab1, tab2 = st.tabs(["CAPTURA DE DATOS", "CRUCE DE INFORMACION"])

    # --- PESTAÑA 1: CAPTURA DE DATOS ---
    with tab1:
        with st.container():
            st.markdown(f"### <img src='{icono_reporte}' width='35' style='vertical-align: middle; margin-right: 10px;'> CAPTURA DE REPORTE MANUAL", unsafe_allow_html=True)
            st.markdown("Ingresa la información de los componentes. Puedes agregar nuevas filas al final de la tabla.")
            
            st.markdown('<div class="marcador-tabla"></div>', unsafe_allow_html=True)
            
            columnas_plantilla = [
                "FECHA", "DESCRIPCIÓN", "NUMERO DE PARTE", "NUMERO DE SERIE", 
                "NUMERO DE PARTE2", "NUMERO DE SERIE3", "EST", "AC", "DOCS",
                "N°TO", "RO/EX", "REGIMEN", "OBSERVACIONES"
            ]
            
            if "df_input" not in st.session_state:
                df_vacio = pd.DataFrame(columns=columnas_plantilla)
                for i in range(5):
                    df_vacio.loc[i] = [None] * 13 # Ahora son 13 columnas
                
                df_vacio["FECHA"] = pd.to_datetime(df_vacio["FECHA"])
                
                for col in columnas_plantilla:
                    if col != "FECHA":
                        df_vacio[col] = df_vacio[col].astype(str)

                df_vacio.index = range(1, len(df_vacio) + 1)
                df_vacio.index.name = "No."

                st.session_state["df_input"] = df_vacio

            edited_df = st.data_editor(
                st.session_state["df_input"],
                num_rows="dynamic",
                use_container_width=True,
                column_config={
                    "FECHA": st.column_config.DateColumn("FECHA", format="DD/MM/YYYY"),
                    "DESCRIPCIÓN": st.column_config.TextColumn("DESCRIPCIÓN"),
                    "NUMERO DE PARTE": st.column_config.TextColumn("INSTALADO: NO. DE PARTE"),
                    "NUMERO DE SERIE": st.column_config.TextColumn("INSTALADO: NO. DE SERIE"),
                    "NUMERO DE PARTE2": st.column_config.TextColumn("REMOVIDO: NO. DE PARTE"),
                    "NUMERO DE SERIE3": st.column_config.TextColumn("REMOVIDO: NO. DE SERIE")
                }
            )

        st.markdown("---")
        
        # 4. Magia final: Convertimos el índice visual de nuevo en la columna "No." para el sistema
        df_final_captura = edited_df.reset_index()
        
        output_captura = BytesIO()
        with pd.ExcelWriter(output_captura, engine='openpyxl') as writer:
            df_final_captura.to_excel(writer, index=False, sheet_name='Reporte_Manual')
        output_captura.seek(0)

        st.download_button(
            label="Descargar Reporte para Enviar",
            data=output_captura,
            file_name="Reporte_Manual_Capturado.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary"
        )

        # Guardamos en memoria la tabla con la columna "No." lista para la Pestaña 2
        st.session_state["datos_capturados"] = df_final_captura.copy()

    # --- PESTAÑA 2: CRUCE DE INFORMACIÓN ---
    with tab2:
        with st.container():
            st.markdown('<div class="marcador-cruce"></div>', unsafe_allow_html=True)
            st.markdown(f"### <img src='{icono_cruce}' width='35' style='vertical-align: middle; margin-right: 10px;'> CRUCE DE INFORMACION", unsafe_allow_html=True)
            st.markdown("Sube el reporte manual capturado y el archivo en crudo descargado del sistema.")
            usar_captura = st.toggle("Usar la tabla recien capturada",value=True)

            df_plantilla_raw = None
            file_descargado = None
            col1, col2 = st.columns(2)
            with col1:
                if usar_captura:
                    st.success("Datos manuales enlazados correctamente desde la Pestaña 1.")
                    # Rescatamos la tabla de la memoria y limpiamos filas vacías
                    if "datos_capturados" in st.session_state:
                        df_plantilla_raw = st.session_state["datos_capturados"].dropna(how='all')
                else:
                    file_plantilla = st.file_uploader("1. Sube el Reporte Manual", type=["xlsx"])
                    if file_plantilla:
                        hojas_p = pd.ExcelFile(file_plantilla).sheet_names
                        hoja_plantilla = st.selectbox("Selecciona la pestaña a usar:", hojas_p) if len(hojas_p) > 1 else hojas_p[0]
                        file_plantilla.seek(0)
                        df_plantilla_raw = cargar_excel(file_plantilla, "plantilla", hoja_plantilla)

            with col2:
                file_descargado = st.file_uploader("2. Sube el Archivo del Sistema", type=["xlsx"])
                if file_descargado:
                    hojas_d = pd.ExcelFile(file_descargado).sheet_names
                    hoja_descargado = st.selectbox("Selecciona la pestaña a usar:", hojas_d) if len(hojas_d) > 1 else hojas_d[0]
                    file_descargado.seek(0)
                    df_descargado = cargar_excel(file_descargado, "sistema", hoja_descargado)
        st.markdown("---")

        if df_plantilla_raw is not None and file_descargado is not None:
            st.info("Procesando cruce de datos...")
            try:
                # Compatibilidad con archivos históricos
                df_plantilla_raw = df_plantilla_raw.rename(columns={
                    "NUMERO DE PARTE.1": "NUMERO DE PARTE2",
                    "NUMERO DE SERIE.1": "NUMERO DE SERIE3"
                })
                
                espacio_alertas = st.container()
                espacio_edicion = st.container()
                espacio_resultados = st.container()
                
                with espacio_edicion:
                    st.markdown("### CORRECCIÓN DE DATOS")
                    st.info("Si la alerta de arriba indica que no hubo coincidencia, busca el registro aquí, corrígelo y presiona 'Enter'.")
                    df_plantilla = st.data_editor(df_plantilla_raw, use_container_width=True)
                    st.markdown("---")
                
                if "NUMERO DE SERIE3" in df_plantilla.columns and "SN" in df_descargado.columns:
                    df_plantilla["NUMERO DE SERIE3"] = df_plantilla["NUMERO DE SERIE3"].astype(str).str.strip()
                    df_descargado["SN"] = df_descargado["SN"].astype(str).str.strip()

                    df_cruce = pd.merge(
                        df_plantilla,
                        df_descargado,
                        left_on="NUMERO DE SERIE3",
                        right_on="SN",
                        how="left"
                    )

                    if "PN_DESCRIPTION" in df_cruce.columns and "DESCRIPCIÓN" in df_cruce.columns:
                        df_cruce["DESCRIPCIÓN"] = np.where(
                            df_cruce["PN_DESCRIPTION"].notna() & (df_cruce["PN_DESCRIPTION"] != "nan") & (df_cruce["PN_DESCRIPTION"].str.strip() != ""), 
                            df_cruce["PN_DESCRIPTION"], 
                            df_cruce["DESCRIPCIÓN"]
                        )
                    
                    if "PN" in df_cruce.columns and "NUMERO DE PARTE2" in df_cruce.columns:
                        df_cruce["NUMERO DE PARTE2"] = np.where(
                            df_cruce["PN"].notna() & (df_cruce["PN"] != "nan") & (df_cruce["PN"].str.strip() != ""), 
                            df_cruce["PN"], 
                            df_cruce["NUMERO DE PARTE2"]
                        )

                    columnas_originales = df_plantilla.columns.tolist()
                    df_final = df_cruce[columnas_originales]
                    
                    df_final = df_final.replace(['nan', 'NaN', '#REF!', 'None', 'NaT'], '')
                    df_final = df_final.fillna('')

                    filas_sin_match = df_final[df_final["DESCRIPCIÓN"] == ""]
                    
                    with espacio_alertas:
                        if not filas_sin_match.empty:
                            with st.container():
                                st.markdown('<div class="marcador-alerta"></div>', unsafe_allow_html=True)
                                st.markdown(f"### <img src='{icono_alerta}' width='35' style='vertical-align: middle; margin-right: 10px;'> ALERTAS", unsafe_allow_html=True)
                                st.error(f"No se encontró coincidencia en el sistema para {len(filas_sin_match)} componentes. Revisa la tabla de corrección abajo.")
                                st.dataframe(filas_sin_match)
                        else:
                            st.success("¡Cruce realizado con éxito para todos los registros!")
                        st.markdown("---")
                    
                    with espacio_resultados:
                        # Truco visual para el Excel de salida
                        df_final = df_final.rename(columns={
                            "NUMERO DE PARTE2": "NUMERO DE PARTE ", 
                            "NUMERO DE SERIE3": "NUMERO DE SERIE "
                        })

                        st.markdown("### Tabla Completa")
                        st.dataframe(df_final)

                        # Títulos superiores del Excel
                        nivel_superior = [
                            "DATOS GENERALES", 
                            "DATOS GENERALES", "DATOS GENERALES", 
                            "COMPONENTE INSTALADO", "COMPONENTE INSTALADO",
                            "COMPONENTE REMOVIDO", "COMPONENTE REMOVIDO",  
                            "DATOS GENERALES", "DATOS GENERALES", "DATOS GENERALES", 
                            "DATOS GENERALES", "DATOS GENERALES", "DATOS GENERALES", "DATOS GENERALES" 
                        ]
                        
                        output_final = BytesIO()
                        with pd.ExcelWriter(output_final, engine='openpyxl') as writer:
                            df_final.to_excel(writer, index=False, sheet_name='Reporte_Completado', startrow=3)
                            
                            workbook = writer.book
                            worksheet = writer.sheets['Reporte_Completado']
                            
                            worksheet['A1'] = "REPORTE DE COMPONENTES ROTABLES"
                            worksheet['A1'].font = Font(size=16, bold=True, color="FFFFFF")
                            worksheet['A1'].alignment = Alignment(horizontal="center", vertical="center")
                            worksheet['A1'].fill = openpyxl.styles.PatternFill(start_color="002060", end_color="002060", fill_type="solid")
                            
                            ultima_col_letra = get_column_letter(len(df_final.columns))
                            worksheet.merge_cells(f'A1:{ultima_col_letra}1')
                            
                            for col_num, valor in enumerate(nivel_superior, 1):
                                celda = worksheet.cell(row=3, column=col_num)
                                celda.value = valor
                                celda.font = Font(bold=True)
                                celda.alignment = Alignment(horizontal="center", vertical="center")
                                celda.fill = openpyxl.styles.PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
                            
                            ultima_fila = len(df_final) + 4 
                            rango_tabla = f"A4:{ultima_col_letra}{ultima_fila}"
                            
                            tabla = Table(displayName="TablaRotables", ref=rango_tabla)
                            estilo = TableStyleInfo(
                                name="TableStyleMedium9", 
                                showFirstColumn=False,
                                showLastColumn=False, 
                                showRowStripes=True, 
                                showColumnStripes=False
                            )
                            tabla.tableStyleInfo = estilo
                            worksheet.add_table(tabla)
                            
                            for col_idx in range(1, worksheet.max_column + 1):
                                column_letter = get_column_letter(col_idx)
                                max_length = 0
                                for row_idx in range(3, worksheet.max_row + 1):
                                    cell = worksheet.cell(row=row_idx, column=col_idx)
                                    try:
                                        if cell.value:
                                            longitud = len(str(cell.value))
                                            if longitud > max_length:
                                                max_length = longitud
                                    except:
                                        pass
                                worksheet.column_dimensions[column_letter].width = max(max_length + 2, 10)

                        output_final.seek(0)

                        st.download_button(
                            label="Descargar Reporte Finalizado",
                            data=output_final,
                            file_name="Reporte_Rotables_Completado.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            type="primary"
                        )
                else:
                    st.error("No se encontraron las columnas necesarias para el cruce. Verifica los archivos subidos.")
            except Exception as e:
                st.error(f"Error procesando los datos: {e}")
